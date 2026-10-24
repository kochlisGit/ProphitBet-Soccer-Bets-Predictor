import os
import pandas as pd
from typing import Optional


def export_table(df: pd.DataFrame, filepath: str, append: bool = True) -> Optional[str]:
    """ Exports a dataframe to a CSV or Excel file.

        Args
        ----
        :param df: (pd.DataFrame) dataframe to be expected as CSV/Excel file.
        :param filepath: (str) path to file to be exported. It supports only xlsx and csv formats.
        :param append: (bool) append to an existing file. If False, replace it.
    """

    split_path = filepath.rsplit('.', maxsplit=1)

    if len(split_path) == 1:
        filepath += '.csv'
        excel = False
    else:
        filetype = split_path[1]

        if filetype == 'csv':
            excel = False
        elif filetype == 'xlsx':
            excel = True
        else:
            return f'Invalid file type: "{filetype}"'

    try:
        if excel:
            if append and os.path.exists(filepath):
                from openpyxl import load_workbook

                workbook = load_workbook(filename=filepath)
                startrow, mode, header, if_sheet_exists = (workbook['Sheet1'].max_row, 'a', False, 'overlay')
            else:
                startrow, mode, header, if_sheet_exists = (0, 'w', True, None)

            with pd.ExcelWriter(path=filepath, engine='openpyxl', mode=mode, if_sheet_exists=if_sheet_exists) as writer:
                df.to_excel(
                    excel_writer=writer,
                    sheet_name='Sheet1',
                    startrow=startrow,
                    header=header,
                    index=False
                )
        else:
            mode, header = ('a', False) if append and os.path.exists(path=filepath) else ('w', True)
            df.to_csv(filepath, mode=mode, header=header, index=False)

        return None
    except Exception as e:
        return f'Could not export data: {e}'
